"""Коммерческое предложение в DOCX по бренду РТК (skill rt-brand, раздел «Шаблон КП»).

Логотип — оригинальный PNG (docs/brand/identity, уменьшен до 1400 px в backend/templates/rt_logo.png), заголовки таблиц
фиолетовые #7700FF, итоговая сумма оранжевая #FF4F12, шрифт Arial (разрешённый гайдлайном фолбэк Rostelecom Basis:
у получателя фирменного шрифта, скорее всего, нет). Под суммой — что не входит в стоимость.

Цены в спецификации — клиентские: внутренние суммы (с коэффициентами сложности) пропорционально приведены к цене
предложения, поэтому сумма строк равна итогу, а внутренняя себестоимость и наценка в документ не попадают.
"""
import io
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LOGO = Path(__file__).resolve().parents[2] / "templates" / "rt_logo.png"
PURPLE, ORANGE, TEXT, MUTED = RGBColor(0x77, 0x00, 0xFF), RGBColor(0xFF, 0x4F, 0x12), RGBColor(0x10, 0x18, 0x28), RGBColor(0x57, 0x5D, 0x68)
FONT = "Arial"
VALID_DAYS = 30

NOT_INCLUDED = [
    "общестроительные и электромонтажные работы, прокладка кабельных трасс (кроме шеф-монтажа)",
    "доработка или замена существующего инженерного оборудования здания",
    "каналы связи для передачи данных — предложим отдельно: Интернет для бизнеса или IoT-сеть Ростелекома",
    "лицензии и абонентская плата за платформу диспетчеризации после ввода в эксплуатацию",
    "сервисное обслуживание системы после окончания гарантийного срока",
]


def _rub(v: float) -> str:
    return f"{round(v):,}".replace(",", " ") + " ₽"


def _plural(n: int, one: str, few: str, many: str) -> str:
    m10, m100 = n % 10, n % 100
    return one if m10 == 1 and m100 != 11 else few if 2 <= m10 <= 4 and not 12 <= m100 <= 14 else many


def _run(p, text, size=10, bold=False, color=TEXT):
    r = p.add_run(text)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return r


def _para(doc, text="", size=10, bold=False, color=TEXT, space_after=4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        _run(p, text, size, bold, color)
    return p


def _shade(cell, hex_fill: str):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _cell(cell, text, size=9, bold=False, color=TEXT, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    _run(p, text, size, bold, color)


def build(req: dict, e: dict, ref: dict, offer: float, object_types: list[dict], manager: dict | None = None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10)

    client = req.get("client") or {}
    today = date.today()
    number = f"КП-{today:%Y%m%d}-{(client.get('inn') or '0000')[-4:]}"
    level = next(l for l in ref["levels"] if l["code"] == e["params"]["level"])
    otype = next((t["label"] for t in object_types if t["code"] == req.get("object_type")), None)

    doc.add_picture(str(LOGO), width=Cm(5))
    _para(doc, "ПАО «Ростелеком» · решения для бизнеса", 8, color=MUTED, space_after=14)
    _para(doc, "Коммерческое предложение", 20, bold=True, space_after=2)
    _para(doc, f"№ {number} от {today:%d.%m.%Y} · действительно до {today + timedelta(days=VALID_DAYS):%d.%m.%Y}", 9, color=MUTED, space_after=12)

    facts = [("Кому", client.get("name") or "—"), ("ИНН", client.get("inn") or "—"),
             ("Объект", ", ".join(x for x in [client.get("object_name"), otype,
                                              f"{req['area_m2']:,.0f} м²".replace(",", " ") if req.get("area_m2") else None] if x) or "—"),
             ("Адрес", client.get("address") or "—")]
    if client.get("contact"):
        facts.append(("Контакт", client["contact"]))
    t = doc.add_table(rows=0, cols=2)
    for k, v in facts:
        row = t.add_row().cells
        _cell(row[0], k, color=MUTED), _cell(row[1], v)
        row[0].width, row[1].width = Cm(3), Cm(14)

    _para(doc, "", space_after=6)
    _para(doc, "Предлагаем", 13, bold=True, space_after=2)
    _para(doc, f"Систему мониторинга и диспетчеризации инженерных систем здания под ключ, уровень «{level['label']}»: "
               f"{level['note'][0].lower() + level['note'][1:]} Проектирование, поставка оборудования, шеф-монтаж и пусконаладка.",
          10, space_after=10)

    # Спецификация: внутренние суммы × k = клиентские, Σ строк = цена предложения
    k = offer / e["subtotal"] if e["subtotal"] else 0
    lines = [(l["name"], l["unit"], l["qty"], l["sum"]) for l in e["lines"]]
    lines += [(c["name"], "шт.", c["qty"], c["sum"]) for c in e["custom"]]
    lines += [(w["label"], "компл.", 1, w["sum"]) for w in e["works"]]
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [Cm(0.9), Cm(8.1), Cm(1.4), Cm(1.4), Cm(2.6), Cm(2.6)]
    heads = ["№", "Наименование", "Ед.", "Кол-во", "Цена, ₽", "Сумма, ₽"]
    right = WD_ALIGN_PARAGRAPH.RIGHT
    for i, (c, h) in enumerate(zip(table.rows[0].cells, heads)):
        _cell(c, h, 9, True, RGBColor(0xFF, 0xFF, 0xFF), right if i >= 3 else None)
        _shade(c, "7700FF")
    for n, (name, unit, qty, s) in enumerate(lines, 1):
        cells = table.add_row().cells
        amount = s * k
        vals = [str(n), name, unit, str(qty), _rub(amount / qty)[:-2], _rub(amount)[:-2]]
        for i, (c, v) in enumerate(zip(cells, vals)):
            _cell(c, v, 9, align=right if i >= 3 else None)
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w

    vat = offer * e["vat"]
    _para(doc, "", space_after=4)
    for label, value in [("Итого без НДС", offer), (f"НДС {e['vat']:.0%}", vat)]:
        p = _para(doc, align=right, space_after=0)
        _run(p, f"{label}:  ", 10, color=MUTED), _run(p, _rub(value), 10)
    p = _para(doc, align=right, space_after=2)
    _run(p, "Итого с НДС:  ", 12, bold=True), _run(p, _rub(offer + vat), 18, bold=True, color=ORANGE)

    _para(doc, "", space_after=4)
    fin = next(c for c in e["coefficients"] if c["kind"] == "financing")
    urg = next(c for c in e["coefficients"] if c["kind"] == "urgency")
    weeks = f"{e['weeks']} {_plural(e['weeks'], 'неделя', 'недели', 'недель')}"
    terms = [("Срок выполнения", f"{weeks} с даты подписания договора" + ("" if urg["value"] == 1 else f", {urg['label'].lower()}")),
             ("Оплата", fin["label"]), ("Гарантия", "12 месяцев на оборудование и работы"),
             ("Действует", f"{VALID_DAYS} дней, до {today + timedelta(days=VALID_DAYS):%d.%m.%Y}")]
    t = doc.add_table(rows=0, cols=2)
    for kk, v in terms:
        row = t.add_row().cells
        _cell(row[0], kk, color=MUTED), _cell(row[1], v)
        row[0].width, row[1].width = Cm(4), Cm(13)

    _para(doc, "", space_after=4)
    _para(doc, "Не входит в стоимость", 10, bold=True, space_after=2)
    for x in NOT_INCLUDED:
        _para(doc, f"— {x}", 9, color=MUTED, space_after=1)

    _para(doc, "", space_after=10)
    is_guest = not manager or manager.get("role") == "guest" or manager.get("email") == "guest" or manager.get("full_name") == "Гость"
    if not is_guest and manager:
        pos = manager.get("position") or "Менеджер по работе с корпоративными клиентами"
        _para(doc, f"{pos} ПАО «Ростелеком»", 10, space_after=2)
        _para(doc, f"______________________  /  {manager.get('full_name') or '______________________'}", 10, space_after=2)
        contacts = " · ".join(x for x in (manager.get("phone"), manager.get("email")) if x and x != "guest")
        if contacts:
            _para(doc, contacts, 9, color=MUTED, space_after=14)
        else:
            _para(doc, "", space_after=14)
    else:
        _para(doc, "Менеджер по работе с корпоративными клиентами ПАО «Ростелеком»", 10, space_after=2)
        _para(doc, "______________________  /  ______________________", 10, space_after=14)
    _para(doc, "ПРЕДЛОЖЕНИЕ НЕ ЯВЛЯЕТСЯ ПУБЛИЧНОЙ ОФЕРТОЙ. ЦЕНА РАССЧИТАНА ПО ТИПОВОЙ КОМПЛЕКТАЦИИ И УТОЧНЯЕТСЯ ПОСЛЕ "
               "ОБСЛЕДОВАНИЯ ОБЪЕКТА. ПОДГОТОВЛЕНО В КАЛЬКУЛЯТОРЕ СДЕЛОК РТК.", 7, color=MUTED)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
